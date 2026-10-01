"""Builds seo_migration_audit.ipynb from the cell sources below."""
import sys
import nbformat as nbf

INTRO = r"""# SEO Migration Audit — Live vs Staging

Compares every page on the **live** site with the same page on **staging** and flags anything that changed:

| Check | What it verifies |
|---|---|
| HTTP status | Staging page loads (200) and is HTML |
| URL | Same path as live, no unexpected redirects, clean URL format |
| Trailing slash | Same slash convention; the other variant 301-redirects to the right one |
| Meta title / description | Identical text after whitespace/entity normalization |
| Canonical | Present, single, self-referencing, same path as live |
| Robots & indexing | `noindex`/`nofollow` (meta + `X-Robots-Tag`) differences |
| Heading structure | Same H1–H6 outline (level + text, in order), exactly one H1 |
| Image alts | Matched images keep their alt; no missing alts on staging |
| Image names | Same file names (catches renames, `.jpg` → `.webp`, `IMG_1234` style names) |
| Image URL structure | Same image folder path (e.g. `/wp-content/uploads/2024/05/`) |
| Schema | Same JSON-LD / microdata `@type`s and properties, valid JSON |
| Open Graph / Hreflang / Content length | Social tags, language alternates, word count drop > 20% |

## How to use
1. **Runtime → Run all** (or run cells ① → ⑤ one by one).
2. Fill in the form in cell ① (sites, staging login if any) and choose where the URL list comes from in cell ②.
3. Cell ⑤ downloads an Excel report: **Overview**, **Summary** (one row per page), **All issues**, one sheet per check, and **Images**.

> ⚠️ Colab runs on Google's servers. If staging only allows your office/VPN IP, ask the developers to allow access (or give a username/password) — otherwise every staging page returns 401/403.
"""

SETTINGS = r'''#@title ① Settings { display-mode: "form" }
LIVE_BASE = "https://transcure.net"  #@param {type:"string"}
STAGE_BASE = "https://stage.transcure.net"  #@param {type:"string"}

#@markdown **Staging access** (leave blank if staging is public)
STAGE_USERNAME = ""  #@param {type:"string"}
STAGE_PASSWORD = ""  #@param {type:"string"}
STAGE_COOKIE = ""  #@param {type:"string"}

#@markdown **Crawl options**
MAX_URLS = 0  #@param {type:"integer"}
#@markdown ↑ 0 = audit every URL. Use e.g. 20 for a quick test run first.
EXCLUDE_REGEX = ""  #@param {type:"string"}
#@markdown ↑ Skip URLs matching this regex, e.g. `/tag/|/author/|/page/\d+`
CHECK_SLASH_REDIRECTS = True  #@param {type:"boolean"}
MAX_WORKERS = 4  #@param {type:"slider", min:1, max:16, step:1}
DELAY_SECONDS = 0.2  #@param {type:"number"}
TIMEOUT = 30  #@param {type:"integer"}
USER_AGENT = "Mozilla/5.0 (compatible; SEO-Migration-Audit/1.0)"  #@param {type:"string"}
'''

URLS = r'''#@title ② URL list { display-mode: "form" }
URL_SOURCE = "Sitemap of live site"  #@param ["Sitemap of live site", "Paste URLs (below)", "Upload CSV"]
SITEMAP_URL = ""  #@param {type:"string"}
#@markdown ↑ Optional. Blank = auto-detect (`robots.txt`, `/sitemap_index.xml`, `/sitemap.xml`).

#@markdown ---
#@markdown **Paste URLs** — used when *Paste URLs* is selected. One per line, either:
#@markdown - `https://transcure.net/page/` (staging URL is generated automatically), or
#@markdown - `https://transcure.net/old-page/, https://stage.transcure.net/new-page/` (for pages whose URL changed)
#@markdown
#@markdown **Upload CSV** — columns `live_url` and (optional) `stage_url`.
PASTED_URLS = """
https://transcure.net/
"""
'''

ENGINE = r'''#@title ③ Audit engine — just run this cell, no edits needed { display-mode: "form" }
import re, json, html as htmllib, time, difflib, posixpath, threading, datetime, io
from urllib.parse import urlsplit, urlunsplit, urljoin, unquote
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from bs4 import BeautifulSoup
import pandas as pd
from tqdm.auto import tqdm

CHECKS = [
    "HTTP status", "URL", "Trailing slash", "Meta title", "Meta description",
    "Canonical", "Robots & indexing", "Heading structure", "Image alts",
    "Image names", "Image URL structure", "Schema", "Open Graph", "Hreflang",
    "Content length",
]
MAX_LIST = 25  # max items listed per finding cell


def _split_base(u):
    p = urlsplit(u.strip().rstrip("/"))
    return (p.scheme or "https"), p.netloc.lower()


LIVE_SCHEME, LIVE_HOST = _split_base(LIVE_BASE)
STAGE_SCHEME, STAGE_HOST = _split_base(STAGE_BASE)


def _host_variants(h):
    bare = h[4:] if h.startswith("www.") else h
    return {bare, "www." + bare}


SITE_HOSTS = _host_variants(LIVE_HOST) | _host_variants(STAGE_HOST)


def strip_host(u):
    """URL -> path(+query) for our own hosts, so live and staging compare equal."""
    if not u:
        return ""
    p = urlsplit(u.strip())
    if p.netloc and p.netloc.lower() not in SITE_HOSTS:
        return urlunsplit((p.scheme, p.netloc, p.path, p.query, ""))
    return (p.path or "/") + ("?" + p.query if p.query else "")


def host_of(u):
    return urlsplit(u).netloc.lower() if u else ""


def norm(s):
    if s is None:
        return ""
    s = htmllib.unescape(str(s)).replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def to_stage(live_url):
    p = urlsplit(live_url)
    return urlunsplit((STAGE_SCHEME, STAGE_HOST, p.path or "/", p.query, ""))


def short(items, n=MAX_LIST):
    items = list(items)
    extra = f"\n… and {len(items) - n} more" if len(items) > n else ""
    return "\n".join(str(i) for i in items[:n]) + extra


# ---------------------------------------------------------------- fetching
_tls = threading.local()


def session(stage):
    key = "stage" if stage else "live"
    s = getattr(_tls, key, None)
    if s is None:
        s = requests.Session()
        s.headers.update({"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml,*/*;q=0.8"})
        retry = Retry(total=2, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504],
                      allowed_methods=["GET"], raise_on_status=False)
        s.mount("http://", HTTPAdapter(max_retries=retry))
        s.mount("https://", HTTPAdapter(max_retries=retry))
        if stage and STAGE_USERNAME:
            s.auth = (STAGE_USERNAME, STAGE_PASSWORD)
        if stage and STAGE_COOKIE:
            s.headers["Cookie"] = STAGE_COOKIE
        setattr(_tls, key, s)
    return s


def _decode(r):
    if "charset" in r.headers.get("Content-Type", "").lower():
        return r.text
    try:
        return r.content.decode("utf-8")
    except UnicodeDecodeError:
        return r.text


def fetch(url, stage, follow=True):
    out = {"url": url, "status": None, "final_url": url, "chain": [], "html": "",
           "ctype": "", "x_robots": "", "location": "", "error": ""}
    try:
        if DELAY_SECONDS:
            time.sleep(DELAY_SECONDS)
        r = session(stage).get(url, timeout=TIMEOUT, allow_redirects=follow)
        out["status"] = r.status_code
        out["final_url"] = r.url
        out["chain"] = [f"{h.status_code} {h.url}" for h in r.history]
        out["ctype"] = r.headers.get("Content-Type", "")
        out["x_robots"] = r.headers.get("X-Robots-Tag", "")
        if r.headers.get("Location"):
            out["location"] = urljoin(r.url, r.headers["Location"])
        if follow and "html" in out["ctype"].lower():
            out["html"] = _decode(r)
        r.close()
    except requests.RequestException as e:
        out["error"] = f"{type(e).__name__}: {e}"[:300]
    return out


def alt_variant(url):
    """The same URL with the trailing slash toggled (None for root/files/query URLs)."""
    p = urlsplit(url)
    path = p.path or "/"
    if path == "/" or p.query or "." in path.rstrip("/").rsplit("/", 1)[-1]:
        return None
    new = path[:-1] if path.endswith("/") else path + "/"
    return urlunsplit((p.scheme, p.netloc, new, "", ""))


# -------------------------------------------------------------- extraction
SIZE_SUFFIX = re.compile(r"(-\d+x\d+|-scaled|-e\d{10,}|@\dx)$", re.I)


def image_info(src, alt):
    path = urlsplit(src).path
    name = unquote(posixpath.basename(path))
    stem, ext = posixpath.splitext(name)
    key = stem.lower()
    while SIZE_SUFFIX.search(key):
        key = SIZE_SUFFIX.sub("", key)
    return {"src": src, "alt": alt, "name": name, "ext": ext.lower(), "key": key,
            "dir": unquote(posixpath.dirname(path)) + "/", "host": host_of(src)}


def _walk_schema(node, acc):
    if isinstance(node, list):
        for n in node:
            _walk_schema(n, acc)
    elif isinstance(node, dict):
        t = node.get("@type")
        if t:
            for tt in (t if isinstance(t, list) else [t]):
                acc.setdefault(str(tt), set()).update(k for k in node if not k.startswith("@"))
        for k, v in node.items():
            if isinstance(v, (dict, list)):
                _walk_schema(v, acc)


def _parse_jsonld(raw):
    raw = re.sub(r"^\s*(<!--|//\s*<!\[CDATA\[)|(-->|//\s*\]\]>)\s*$", "", raw.strip())
    try:
        return json.loads(raw)
    except ValueError:
        return json.loads(re.sub(r",\s*([}\]])", r"\1", raw))  # tolerate trailing commas


def extract(html, base):
    soup = BeautifulSoup(html, "lxml")
    head = soup.head or soup
    d = {}

    t = head.find("title") or soup.find("title")
    d["title"] = norm(t.get_text()) if t else ""

    def meta(attr, value):
        tag = soup.find("meta", attrs={attr: re.compile(rf"^{re.escape(value)}$", re.I)})
        return norm(tag.get("content")) if tag and tag.get("content") is not None else ""

    d["description"] = meta("name", "description")
    d["robots"] = ", ".join(x for x in (meta("name", "robots"), meta("name", "googlebot")) if x)
    d["og"] = {k: meta("property", k) for k in ("og:title", "og:description", "og:image")}

    d["canonicals"], d["hreflang"] = [], {}
    for link in soup.find_all("link", href=True):
        rels = [r.lower() for r in (link.get("rel") or [])]
        href = urljoin(base, link["href"].strip())
        if "canonical" in rels:
            d["canonicals"].append(href)
        if "alternate" in rels and link.get("hreflang"):
            d["hreflang"][link["hreflang"].lower()] = strip_host(href)

    d["headings"] = [(int(h.name[1]), norm(h.get_text(" ", strip=True)))
                     for h in soup.find_all(re.compile(r"^h[1-6]$"))]

    imgs, seen = [], set()
    for img in soup.find_all("img"):
        src = None
        for a in ("data-src", "data-lazy-src", "data-original", "src"):
            v = (img.get(a) or "").strip()
            if v and not v.startswith("data:"):
                src = v
                break
        if not src:
            srcset = (img.get("data-srcset") or img.get("srcset") or "").strip()
            src = srcset.split(",")[0].strip().split(" ")[0] if srcset else None
        if not src:
            continue
        src = urljoin(base, src)
        if src in seen:
            continue
        seen.add(src)
        imgs.append(image_info(src, img.get("alt")))
    d["images"] = imgs

    schema, errors, d["schema_has_stage_urls"] = {}, [], False
    for s in soup.find_all("script", type=re.compile(r"ld\+json", re.I)):
        raw = s.string or s.get_text() or ""
        if not raw.strip():
            continue
        d["schema_has_stage_urls"] |= STAGE_HOST in raw
        try:
            _walk_schema(_parse_jsonld(raw), schema)
        except ValueError as e:
            errors.append(f"Invalid JSON-LD: {str(e)[:120]}")
    for el in soup.find_all(attrs={"itemtype": True}):
        for it in el["itemtype"].split():
            schema.setdefault(it.rstrip("/").rsplit("/", 1)[-1], set())
    d["schema"], d["schema_errors"] = schema, errors

    for tag in soup(["script", "style", "noscript", "svg", "template"]):
        tag.decompose()
    body = soup.body or soup
    d["words"] = len(re.findall(r"\w+", body.get_text(" ")))
    return d


# -------------------------------------------------------------- comparison
def R(check, status, live="", stage="", note=""):
    return {"check": check, "status": status, "live": live, "stage": stage, "note": note}


def _url_hygiene(path):
    issues = []
    p = path.split("?")[0]
    if re.search(r"[A-Z]", p): issues.append("uppercase letters")
    if "_" in p: issues.append("underscores")
    if re.search(r"%20|\s", p): issues.append("spaces")
    if "//" in p: issues.append("double slash")
    return issues


def _is_noindex(s):
    return "noindex" in s.lower() or "none" in [x.strip() for x in s.lower().split(",")]


BAD_IMG_NAME = re.compile(r"^(img|dsc|dcim|image|screenshot|photo|untitled|pxl|wp)[-_ ]?\d+|^[0-9a-f]{16,}$|^\d+$", re.I)


def compare(live, stage, mapped, live_alt, stage_alt):
    out = []
    # HTTP status
    if stage["error"]:
        out.append(R("HTTP status", "FAIL", live["status"], stage["error"], "Staging request failed"))
        return out
    if stage["status"] != 200:
        hint = " — staging login/IP allowlist needed?" if stage["status"] in (401, 403) else ""
        out.append(R("HTTP status", "FAIL", live["status"], stage["status"], f"Staging page does not return 200{hint}"))
        return out
    if not stage["html"]:
        out.append(R("HTTP status", "FAIL", live["status"], stage["ctype"], "Staging response is not HTML"))
        return out
    if live["error"] or live["status"] != 200 or not live["html"]:
        out.append(R("HTTP status", "WARN", live["error"] or live["status"], 200,
                     "Live page itself is not a 200 HTML page — only staging checks were run"))
        live_ok = False
    else:
        out.append(R("HTTP status", "PASS", 200, 200))
        live_ok = True

    L = extract(live["html"], live["final_url"]) if live_ok else None
    S = extract(stage["html"], stage["final_url"])
    live_path, stage_path = strip_host(live["final_url"]), strip_host(stage["final_url"])
    req_path = strip_host(stage["url"])

    # URL
    hyg = _url_hygiene(stage_path)
    chain = " → ".join(stage["chain"])
    if stage["chain"] and stage_path != req_path:
        out.append(R("URL", "FAIL", live_path, stage_path, f"Staging redirects: {chain}"))
    elif mapped:
        out.append(R("URL", "WARN", strip_host(live["url"]), stage_path,
                     "URL changed in migration — make sure the old URL 301-redirects to the new one at launch"))
    elif live_ok and live_path != stage_path:
        out.append(R("URL", "FAIL", live_path, stage_path, "Final path differs from live"))
    elif hyg:
        out.append(R("URL", "WARN", live_path, stage_path, "URL contains " + ", ".join(hyg)))
    else:
        out.append(R("URL", "PASS", live_path, stage_path))

    # Trailing slash
    sp = stage_path.split("?")[0]
    lp = live_path.split("?")[0]
    if sp != "/" and "." not in sp.rstrip("/").rsplit("/", 1)[-1]:
        notes, status = [], "PASS"
        if live_ok and lp.endswith("/") != sp.endswith("/"):
            status = "FAIL"
            notes.append(f"Live URL {'ends' if lp.endswith('/') else 'does not end'} with a slash, staging {'does' if sp.endswith('/') else 'does not'}")

        def describe(a):
            if a is None:
                return ""
            if a["error"]:
                return f"{a['url']} → error"
            return f"{a['url']} → {a['status']}" + (f" {strip_host(a['location'])}" if a["location"] else "")

        if stage_alt is not None and not stage_alt["error"]:
            st, loc = stage_alt["status"], strip_host(stage_alt["location"])
            live_redirects = bool(live_alt and live_alt["status"] in (301, 308))
            if st in (301, 308) and loc == stage_path:
                pass
            elif st in (301, 302, 307, 308) and loc != stage_path:
                status = "FAIL"
                notes.append(f"Other slash variant redirects to the wrong URL ({loc})")
            elif st in (302, 307):
                status = "WARN" if status == "PASS" else status
                notes.append("Other slash variant uses a temporary redirect (302/307) — should be 301")
            elif st == 200:
                status = "FAIL"
                notes.append("Both slash variants return 200 (duplicate content) — the other variant should 301")
            elif live_redirects:
                status = "WARN" if status == "PASS" else status
                notes.append(f"Other slash variant returns {st} on staging but 301s on live (links to it would break)")
        out.append(R("Trailing slash", status, describe(live_alt) or lp, describe(stage_alt) or sp, "; ".join(notes)))

    # Titles & descriptions
    for check, key, limit in (("Meta title", "title", 60), ("Meta description", "description", 160)):
        lv, sv = (L[key] if L else ""), S[key]
        length = f"Staging length {len(sv)}" + (f" (> {limit})" if len(sv) > limit else "")
        if not sv:
            out.append(R(check, "FAIL", lv, sv, "Missing on staging"))
        elif not L:
            out.append(R(check, "PASS", lv, sv, length))
        elif lv == sv:
            out.append(R(check, "PASS", lv, sv, length))
        elif lv.lower() == sv.lower():
            out.append(R(check, "WARN", lv, sv, "Only letter case differs"))
        else:
            out.append(R(check, "FAIL", lv, sv, "Text differs from live" if lv else "Live has none, staging added one"))

    # Canonical
    lc = L["canonicals"] if L else []
    sc = S["canonicals"]
    if not sc:
        out.append(R("Canonical", "FAIL" if lc else "WARN", short(lc), "", "Missing on staging"))
    else:
        expected = strip_host(lc[0]) if lc else stage_path
        got = strip_host(sc[0])
        h = host_of(sc[0])
        if len(sc) > 1:
            out.append(R("Canonical", "FAIL", short(lc), short(sc), "Multiple canonical tags on staging"))
        elif h not in SITE_HOSTS:
            out.append(R("Canonical", "FAIL", short(lc), sc[0], f"Canonical points to another domain ({h})"))
        elif got != expected:
            out.append(R("Canonical", "FAIL", short(lc), sc[0], f"Canonical path {got} ≠ expected {expected}"))
        elif got != stage_path:
            out.append(R("Canonical", "WARN", short(lc), sc[0], "Not self-referencing (same as live, though)"))
        else:
            note = "Uses staging domain — confirm it switches to production at launch" if h in _host_variants(STAGE_HOST) else ""
            out.append(R("Canonical", "PASS", short(lc), sc[0], note))

    # Robots
    lr = ", ".join(x for x in ((L["robots"] if L else ""), live["x_robots"] if live_ok else "") if x)
    sr = ", ".join(x for x in (S["robots"], stage["x_robots"]) if x)
    if _is_noindex(sr) and not _is_noindex(lr):
        out.append(R("Robots & indexing", "WARN", lr or "(none)", sr,
                     "noindex on staging — fine for staging, MUST be removed at launch"))
    elif _is_noindex(lr) != _is_noindex(sr) or ("nofollow" in lr.lower()) != ("nofollow" in sr.lower()):
        out.append(R("Robots & indexing", "FAIL", lr or "(none)", sr or "(none)", "Indexing directives differ"))
    else:
        out.append(R("Robots & indexing", "PASS", lr or "(none)", sr or "(none)"))

    # Headings
    s_lines = [f"H{l}: {t}" for l, t in S["headings"]]
    h1 = sum(1 for l, _ in S["headings"] if l == 1)
    skips = [f"H{a}→H{b}" for (a, _), (b, _) in zip(S["headings"], S["headings"][1:]) if b > a + 1]
    extra = []
    if h1 != 1:
        extra.append(f"{h1} H1 tags on staging (should be 1)")
    if skips:
        extra.append("Skipped levels: " + ", ".join(sorted(set(skips))))
    if L:
        l_lines = [f"H{l}: {t}" for l, t in L["headings"]]
        if l_lines == s_lines:
            out.append(R("Heading structure", "WARN" if h1 != 1 else "PASS", short(l_lines, 60), short(s_lines, 60), "; ".join(extra)))
        else:
            diff = [d for d in difflib.ndiff(l_lines, s_lines) if d[:1] in "-+"]
            diff = [("Missing on staging → " if d[0] == "-" else "Added on staging  → ") + d[2:] for d in diff]
            out.append(R("Heading structure", "FAIL", short(l_lines, 60), short(s_lines, 60), short(diff + extra)))
    else:
        out.append(R("Heading structure", "WARN" if extra else "PASS", "", short(s_lines, 60), "; ".join(extra)))

    # Images
    image_rows = []
    s_imgs = S["images"]
    l_imgs = L["images"] if L else []
    unused = list(s_imgs)
    pairs = []
    for li in l_imgs:
        m = next((si for si in unused if si["name"].lower() == li["name"].lower()), None) \
            or next((si for si in unused if si["key"] == li["key"]), None)
        if m:
            unused.remove(m)
        pairs.append((li, m))

    alt_issues, name_issues, path_issues = [], [], []
    for li, si in pairs:
        row = {"live_src": li["src"], "stage_src": si["src"] if si else "",
               "live_alt": li["alt"], "stage_alt": si["alt"] if si else None}
        if not si:
            name_issues.append(f"Missing/renamed on staging: {li['name']}")
            row.update(name_match="NOT FOUND", alt_match="", path_match="")
        else:
            row["name_match"] = "YES" if si["name"] == li["name"] else "NO"
            row["alt_match"] = "YES" if norm(si["alt"]) == norm(li["alt"]) else "NO"
            row["path_match"] = "YES" if si["dir"] == li["dir"] else "NO"
            if si["name"] != li["name"]:
                name_issues.append(f"{li['name']} → {si['name']}")
            if norm(si["alt"]) != norm(li["alt"]):
                alt_issues.append(f"{si['name']}: \"{norm(li['alt'])}\" → \"{norm(si['alt'])}\"")
            if si["dir"] != li["dir"]:
                path_issues.append(f"{li['name']}: {li['dir']} → {si['dir']}")
        image_rows.append(row)
    for si in unused:
        image_rows.append({"live_src": "", "stage_src": si["src"], "live_alt": None, "stage_alt": si["alt"],
                           "name_match": "NEW ON STAGING", "alt_match": "", "path_match": ""})

    missing_alt = [si["name"] for si in s_imgs if not norm(si["alt"])]
    bad_names = [si["name"] for si in s_imgs if BAD_IMG_NAME.search(si["key"])]
    n_img = f"{len(l_imgs)} live / {len(s_imgs)} staging images"

    notes = alt_issues + ([f"Missing/empty alt on staging ({len(missing_alt)}): " + ", ".join(missing_alt[:MAX_LIST])] if missing_alt else [])
    out.append(R("Image alts", "FAIL" if alt_issues else ("WARN" if missing_alt else "PASS"), n_img, n_img, short(notes)))

    hard = [n for n in name_issues if n.startswith("Missing")]
    if bad_names:
        name_issues.append("Non-descriptive file names: " + ", ".join(bad_names[:MAX_LIST]))
    if unused and L:
        name_issues.append(f"{len(unused)} image(s) only on staging: " + ", ".join(si["name"] for si in unused[:MAX_LIST]))
    out.append(R("Image names", "FAIL" if hard else ("WARN" if name_issues else "PASS"), n_img, n_img, short(name_issues)))

    l_dirs = sorted({li["dir"] for li in l_imgs})
    s_dirs = sorted({si["dir"] for si in s_imgs})
    ext_hosts = sorted({si["host"] for si in s_imgs if si["host"] and si["host"] not in SITE_HOSTS})
    status = "FAIL" if path_issues else ("WARN" if ext_hosts else "PASS")
    if ext_hosts:
        path_issues.append("Images served from other host(s): " + ", ".join(ext_hosts))
    out.append(R("Image URL structure", status, short(l_dirs), short(s_dirs), short(path_issues)))

    # Schema
    ls = L["schema"] if L else {}
    ss = S["schema"]
    fmt = lambda sch: ", ".join(sorted(sch)) or "(none)"
    issues, status = list(S["schema_errors"]), "FAIL" if S["schema_errors"] else "PASS"
    missing = sorted(set(ls) - set(ss))
    added = sorted(set(ss) - set(ls))
    if missing:
        status = "FAIL"
        issues.append("Types missing on staging: " + ", ".join(missing))
    if added and L:
        issues.append("Types added on staging: " + ", ".join(added))
    for t in sorted(set(ls) & set(ss)):
        miss_props = sorted(ls[t] - ss[t])
        if miss_props:
            status = "WARN" if status == "PASS" else status
            issues.append(f"{t} lost properties: " + ", ".join(miss_props))
    if S["schema_has_stage_urls"]:
        issues.append("Schema contains staging URLs — confirm they switch to production at launch")
    if not ls and not ss:
        issues.append("No schema on either version")
    out.append(R("Schema", status, fmt(ls), fmt(ss), short(issues)))

    # Open Graph
    if L:
        diffs = []
        for k in ("og:title", "og:description", "og:image"):
            a, b = L["og"][k], S["og"][k]
            if k == "og:image":
                a, b = strip_host(a), strip_host(b)
            if a != b:
                diffs.append(f"{k}: \"{a}\" → \"{b}\"")
        out.append(R("Open Graph", "WARN" if diffs else "PASS",
                     short(f"{k}: {v}" for k, v in L["og"].items() if v),
                     short(f"{k}: {v}" for k, v in S["og"].items() if v), short(diffs)))

    # Hreflang
    if (L and L["hreflang"]) or S["hreflang"]:
        lh = L["hreflang"] if L else {}
        sh = S["hreflang"]
        fmt_h = lambda h: short(f"{k}: {v}" for k, v in sorted(h.items()))
        out.append(R("Hreflang", "PASS" if lh == sh else "FAIL", fmt_h(lh), fmt_h(sh),
                     "" if lh == sh else "hreflang set differs"))

    # Content length
    if L:
        lw, sw = L["words"], S["words"]
        drop = (lw - sw) / lw if lw else 0
        out.append(R("Content length", "WARN" if drop > 0.2 else "PASS", lw, sw,
                     f"Staging has {drop:.0%} fewer words — content may be missing" if drop > 0.2 else ""))
    return out, image_rows


def audit_pair(idx, live_url, stage_url, mapped):
    try:
        live = fetch(live_url, stage=False)
        stage = fetch(stage_url, stage=True)
        live_alt = stage_alt = None
        if CHECK_SLASH_REDIRECTS:
            a = alt_variant(live["final_url"]) if not live["error"] else None
            b = alt_variant(stage["final_url"]) if not stage["error"] else None
            live_alt = fetch(a, stage=False, follow=False) if a else None
            stage_alt = fetch(b, stage=True, follow=False) if b else None
        res = compare(live, stage, mapped, live_alt, stage_alt)
        checks, images = res if isinstance(res, tuple) else (res, [])
        stage_status = stage["status"] if not stage["error"] else "ERROR"
    except Exception as e:  # never let one page stop the whole audit
        checks, images, stage_status = [R("HTTP status", "FAIL", "", "", f"Audit error: {type(e).__name__}: {e}")], [], "ERROR"
    for c in checks:
        c.update(idx=idx, live_url=live_url, stage_url=stage_url)
    for im in images:
        im.update(idx=idx, page_live=live_url, page_stage=stage_url)
    return idx, stage_status, checks, images


# ---------------------------------------------------------------- URL list
def _get_xml(url):
    r = session(False).get(url, timeout=TIMEOUT)
    if r.status_code != 200:
        return None
    content = r.content
    if url.endswith(".gz") or content[:2] == b"\x1f\x8b":
        import gzip
        content = gzip.decompress(content)
    return BeautifulSoup(content, "xml")


def sitemap_urls(start=None):
    root = f"{LIVE_SCHEME}://{LIVE_HOST}"
    candidates = [start] if start else []
    if not start:
        try:
            robots = session(False).get(root + "/robots.txt", timeout=TIMEOUT).text
            candidates += re.findall(r"(?im)^\s*sitemap:\s*(\S+)", robots)
        except requests.RequestException:
            pass
        candidates += [root + "/sitemap_index.xml", root + "/sitemap.xml", root + "/wp-sitemap.xml"]
    urls, seen = [], set()

    def crawl(sm, depth=0):
        if sm in seen or depth > 5:
            return
        seen.add(sm)
        soup = _get_xml(sm)
        if soup is None:
            return
        if soup.find("sitemapindex"):
            for loc in soup.select("sitemap > loc"):
                crawl(loc.get_text(strip=True), depth + 1)
        for loc in soup.select("url > loc"):
            urls.append(loc.get_text(strip=True))

    for c in candidates:
        crawl(c)
        if urls:
            print(f"Sitemap: {c}")
            break
    return list(dict.fromkeys(urls))


def build_pairs(rows):
    """rows: list of (live, stage_or_None) -> [(idx, live, stage, mapped)]."""
    pairs, seen = [], set()
    rx = re.compile(EXCLUDE_REGEX) if EXCLUDE_REGEX else None
    for live, stage in rows:
        live = live.strip()
        if not live:
            continue
        if live.startswith("/"):
            live = f"{LIVE_SCHEME}://{LIVE_HOST}{live}"
        if rx and rx.search(live):
            continue
        stage = (stage or "").strip()
        if stage.startswith("/"):
            stage = f"{STAGE_SCHEME}://{STAGE_HOST}{stage}"
        stage = stage or to_stage(live)
        mapped = strip_host(stage) != strip_host(live)
        if (live, stage) in seen:
            continue
        seen.add((live, stage))
        pairs.append((len(pairs), live, stage, mapped))
    return pairs[:MAX_URLS] if MAX_URLS else pairs


def run_audit(pairs):
    summary, details, images = {}, [], []
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futs = [ex.submit(audit_pair, *p) for p in pairs]
        for f in tqdm(as_completed(futs), total=len(futs), desc="Auditing pages"):
            idx, stage_status, checks, imgs = f.result()
            details += checks
            images += imgs
            row = {"live_url": pairs[idx][1], "stage_url": pairs[idx][2], "stage_status": stage_status}
            row.update({c["check"]: c["status"] for c in checks})
            row["FAIL"] = sum(c["status"] == "FAIL" for c in checks)
            row["WARN"] = sum(c["status"] == "WARN" for c in checks)
            summary[idx] = row
    summary_df = pd.DataFrame([summary[i] for i in sorted(summary)])
    for c in CHECKS:
        if c not in summary_df:
            summary_df[c] = ""
    summary_df = summary_df[["live_url", "stage_url", "stage_status", "FAIL", "WARN"] + CHECKS].fillna("—")
    details_df = pd.DataFrame(details, columns=["idx", "live_url", "stage_url", "check", "status", "live", "stage", "note"])
    details_df["order"] = details_df["check"].map({c: i for i, c in enumerate(CHECKS)})
    details_df = details_df.sort_values(["idx", "order"]).drop(columns=["idx", "order"]).reset_index(drop=True)
    images_df = pd.DataFrame(images, columns=["idx", "page_live", "page_stage", "live_src", "stage_src", "live_alt",
                                              "stage_alt", "name_match", "alt_match", "path_match"])
    images_df = images_df.sort_values("idx", kind="stable").drop(columns=["idx"]).reset_index(drop=True)
    overview = (details_df.pivot_table(index="check", columns="status", values="live_url", aggfunc="count", fill_value=0)
                .reindex(CHECKS).fillna(0).astype(int))
    for s in ("FAIL", "WARN", "PASS"):
        if s not in overview:
            overview[s] = 0
    overview = overview[["FAIL", "WARN", "PASS"]].reset_index()
    return overview, summary_df, details_df, images_df


def export_excel(path, overview, summary_df, details_df, images_df):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter

    clip = lambda df: df.map(lambda v: v[:32000] if isinstance(v, str) else v) if hasattr(df, "map") else df.applymap(lambda v: v[:32000] if isinstance(v, str) else v)
    issues = details_df[details_df["status"] != "PASS"]
    sheets = [("Overview", overview), ("Summary", summary_df), ("All issues", issues)]
    for c in CHECKS:
        part = issues[issues["check"] == c]
        if len(part):
            sheets.append((c[:31], part.drop(columns=["check"])))
    sheets += [("Images", images_df), ("All details", details_df)]

    red = PatternFill("solid", start_color="F8D7DA")
    amber = PatternFill("solid", start_color="FFF3CD")
    green = PatternFill("solid", start_color="D4EDDA")
    header = PatternFill("solid", start_color="1F3A5F")
    widths = {"live_url": 45, "stage_url": 45, "page_live": 45, "page_stage": 45, "live": 50, "stage": 50,
              "note": 60, "live_src": 50, "stage_src": 50, "live_alt": 35, "stage_alt": 35, "check": 22}

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets:
            clip(df).to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2" if name == "Overview" else "C2" if name == "Summary" else "B2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=i)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14 if name == "Summary" else 18)
            if ws.max_row > 1:
                rng = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=red))
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"WARN"'], fill=amber))
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=green))
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NO"'], fill=red))
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NOT FOUND"'], fill=red))
                if name not in ("Overview", "Summary"):
                    for row in ws.iter_rows(min_row=2):
                        for cell in row:
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return path


print(f"Engine ready. Live: {LIVE_HOST}  |  Staging: {STAGE_HOST}")
'''

RUN = r'''#@title ④ Run the audit { display-mode: "form" }
if URL_SOURCE.startswith("Sitemap"):
    rows = [(u, None) for u in sitemap_urls(SITEMAP_URL or None)]
    if not rows:
        raise SystemExit("No URLs found in the sitemap. Set SITEMAP_URL in cell ② or switch to 'Paste URLs'.")
elif URL_SOURCE.startswith("Paste"):
    rows = []
    for line in PASTED_URLS.strip().splitlines():
        parts = [p.strip() for p in re.split(r"[,\t]", line) if p.strip()]
        if parts:
            rows.append((parts[0], parts[1] if len(parts) > 1 else None))
else:
    from google.colab import files
    up = files.upload()
    csv = pd.read_csv(io.BytesIO(next(iter(up.values()))))
    csv.columns = [c.strip().lower() for c in csv.columns]
    rows = [(r["live_url"], r.get("stage_url") if isinstance(r.get("stage_url"), str) else None)
            for _, r in csv.iterrows()]

pairs = build_pairs(rows)
print(f"{len(pairs)} page pairs to audit (≈{len(pairs) * (4 if CHECK_SLASH_REDIRECTS else 2)} requests)")
for _, l, s, m in pairs[:3]:
    print(f"  {l}  →  {s}{'  (mapped)' if m else ''}")

t0 = time.time()
overview, summary_df, details_df, images_df = run_audit(pairs)
print(f"\nDone in {time.time() - t0:.0f}s")

def _color(v):
    return {"FAIL": "background-color:#f8d7da", "WARN": "background-color:#fff3cd",
            "PASS": "background-color:#d4edda"}.get(v, "")

display(overview.style.set_caption("Pages per check"))
worst = summary_df.sort_values(["FAIL", "WARN"], ascending=False).head(25)
styler = worst.style.map(_color) if hasattr(worst.style, "map") else worst.style.applymap(_color)
display(styler.set_caption("Pages with the most issues (top 25)"))
'''

EXPORT = r'''#@title ⑤ Download Excel report { display-mode: "form" }
report = f"seo_migration_audit_{datetime.datetime.now():%Y-%m-%d_%H%M}.xlsx"
export_excel(report, overview, summary_df, details_df, images_df)
print(f"Saved {report}")
try:
    from google.colab import files
    files.download(report)
except ImportError:
    pass
'''

EXPLORE = r"""### Tip: look at issues without downloading
Run any of these in a new code cell:
```python
details_df[(details_df.check == "Meta title") & (details_df.status == "FAIL")]
details_df[details_df.live_url == "https://transcure.net/some-page/"]
images_df[images_df.alt_match == "NO"]
```
"""

nb = nbf.v4.new_notebook()
nb.metadata = {"colab": {"provenance": [], "toc_visible": True},
               "kernelspec": {"name": "python3", "display_name": "Python 3"},
               "language_info": {"name": "python"}}
nb.cells = [
    nbf.v4.new_markdown_cell(INTRO),
    nbf.v4.new_code_cell(SETTINGS),
    nbf.v4.new_code_cell(URLS),
    nbf.v4.new_code_cell(ENGINE),
    nbf.v4.new_code_cell(RUN),
    nbf.v4.new_code_cell(EXPORT),
    nbf.v4.new_markdown_cell(EXPLORE),
]
for c in nb.cells:
    c.metadata = {"cellView": "form"} if c.cell_type == "code" else {}
nbf.write(nb, sys.argv[1])
print("wrote", sys.argv[1])
